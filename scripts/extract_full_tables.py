"""Full-fidelity extraction of every DIN sheet.
Each sheet -> list of tables. Two kinds:
  h : sizes across a header row; every source row kept as {sym, q (qualifier), v:[value per size]}
  g : any other table kept as a complete grid: head rows (with col/row spans) + body rows
Every numeric cell of the table area must land in a table (checked by the audit)."""
import sys, re, json, glob, collections, statistics
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import extract_skus as X
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, 'output')

# ISO 261 pitches (coarse first) used to catch impossible P values in the source
PITCH = {'1': [.25, .2], '1.2': [.25, .2], '1.4': [.3, .2], '1.6': [.35, .2], '1.8': [.35, .2], '2': [.4, .25], '2.2': [.45, .25], '2.5': [.45, .35], '3': [.5, .35], '3.5': [.6, .35],
         '4': [.7, .5], '5': [.8, .5], '6': [1, .75, .5], '7': [1, .75], '8': [1.25, 1, .75], '10': [1.5, 1.25, 1, .75], '12': [1.75, 1.5, 1.25, 1], '14': [2, 1.5, 1.25, 1],
         '16': [2, 1.5, 1], '18': [2.5, 2, 1.5, 1], '20': [2.5, 2, 1.5, 1], '22': [2.5, 2, 1.5, 1], '24': [3, 2, 1.5, 1], '27': [3, 2, 1.5, 1], '30': [3.5, 3, 2, 1.5, 1],
         '33': [3.5, 3, 2, 1.5], '36': [4, 3, 2, 1.5], '39': [4, 3, 2, 1.5], '42': [4.5, 4, 3, 2, 1.5], '45': [4.5, 4, 3, 2, 1.5], '48': [5, 4, 3, 2, 1.5], '52': [5, 4, 3, 2, 1.5],
         '56': [5.5, 4, 3, 2, 1.5], '60': [5.5, 4, 3, 2, 1.5], '64': [6, 4, 3, 2, 1.5], '68': [6, 4, 3, 2, 1.5], '72': [6, 4, 3, 2], '76': [6, 4, 3, 2], '80': [6, 4, 3, 2]}

QFA = [  # qualifier translation, longest first
    (r'nominal size of pliers.*', 'اندازهٔ انبر خار (DIN 5254)'),
    (r'production class a min\.?', 'حداقل · کلاس A'), (r'production class a m[ai][xi]\.?', 'حداکثر · کلاس A'),
    (r'production class b min\.?', 'حداقل · کلاس B'), (r'production class b max\.?', 'حداکثر · کلاس B'),
    (r'production class c min\.?', 'حداقل · کلاس C'), (r'production class c max\.?', 'حداکثر · کلاس C'),
    (r'production class', 'کلاس تولید'), (r'nominal max\.?', 'اسمی / حداکثر'), (r'nominal min\.?', 'اسمی / حداقل'), (r'nominal size', 'اندازهٔ اسمی'),
    (r'nominal', 'اسمی'), (r'nom\.?', 'اسمی'), (r'max\.?', 'حداکثر'), (r'min\.?', 'حداقل'), (r'approx\.?|≈', 'تقریبی'),
    (r'per\.?\s*dev\.?', 'تلرانس'), (r'tolerance', 'تلرانس'), (r'weight\s*kg\s*/\s*1000\s*pcs\.?|weight.*1000', 'وزن (کیلوگرم در ۱۰۰۰ عدد)'),
    (r'weight', 'وزن'), (r'^kn$', 'kN'), (r'^min-1$', 'دور بر دقیقه'), (r'for thread', 'برای رزوه'), (r'^size$', 'سایز'),
    (r'shaft diameter', 'قطر محور'), (r'bore diameter', 'قطر سوراخ'), (r'^clip$', 'رینگ'), (r'^groove$', 'شیار'),
    (r'supplementary', 'اطلاعات تکمیلی'), (r'^data$', ''), (r'thread', 'رزوه'), (r'závit\s*d', 'رزوه d'), (r'^závit$', 'رزوه'), (r'length', 'طول'),
    (r'for size', 'برای سایز'), (r'pliers', 'انبر خار'), (r'nonstandard dimensions', 'ابعاد غیرمتداول'),
]
def qfa(t):
    s = t.strip().lower()
    if not s:
        return ''
    for pat, fa in QFA:
        if re.fullmatch(pat, s) or (len(pat) > 12 and re.search(pat, s)):
            return fa
    return ''

def clean(v):
    """cell -> display text with '.' decimals; keeps tolerances/tokens."""
    t = X.txt(v)
    if isinstance(v, float):
        t = ('%g' % v)
    t = re.sub(r'(?<=\d),(?=\d)', '.', t)
    return t

def isnum(t):
    return bool(re.fullmatch(r'[+\-−]?\(?\d+(?:\.\d+)?\)?', t.strip()))

def num(t):
    try:
        return float(t.strip('()+').replace('−', '-'))
    except Exception:
        return None

def load_parents():
    wb = openpyxl.load_workbook(ROOT + '/output/DIN_SKU_list.xlsx', read_only=True)
    P = list(wb['Parents'].iter_rows(values_only=True)); hp = P[0]
    return {r[hp.index('Parent code')]: dict(zip(hp, r)) for r in P[1:]}

DUP = [0]
def sheet_tables(sh, ws):
    fixes = []
    for k, v in list(sh.raw.items()):
        t = X.txt(v)
        m = re.match(r'^(\(\s*M[\d,\.]+(?:x[\d,\.]+)?\s*\))\d$', t)
        if m: sh.raw[k] = m.group(1)
        m = re.match(r'^(M\d+)x,(\d)(\d)$', t)  # typo "M30x,15" -> "M30x1,5"
        if m: sh.raw[k] = '%sx%s,%s' % m.groups()
    area = {k: v for k, v in sh.raw.items() if k not in sh.info_cells and k[1] < sh.info_col}
    isint = {k for k, v in area.items() if isinstance(v, int) or (isinstance(v, float) and v.is_integer() and abs(v) >= 100)}
    rows = sorted({r for r, c in area})
    # blocks separated by empty rows; header-only blocks join the next block
    blocks, cur, last = [], [], None
    CAP = re.compile(r'^(tab|table)\.?\s*\d+\.?$', re.I)
    caps, pend = {}, None
    for r in rows:
        first = X.txt(area[min((k for k in area if k[0] == r), key=lambda k: k[1])])
        is_cap = bool(CAP.match(first))
        if (last is not None and r - last > 1) or (is_cap and cur):
            blocks.append(cur); cur = []
        if is_cap:
            pend = first; last = r; continue
        if pend and not cur:
            caps[r] = pend; pend = None
        cur.append(r); last = r
    if cur: blocks.append(cur)
    merged = []
    for b in blocks:
        has_num = any(isnum(clean(area[(r, c)])) for r in b for c in range(1, sh.ncol + 1) if (r, c) in area)
        if merged and not merged[-1][1]:
            merged[-1][0].extend(b); merged[-1][1] = has_num
        else:
            merged.append([list(b), has_num])
    tables = []
    for b, _ in merged:
        cols = sorted({c for (r, c) in area if r in b})
        if not cols:
            continue
        n0 = len(tables)
        # size header rows: >=2 thread tokens (or a 'size' label with numbers)
        size_rows = []
        for r in b:
            cells = [(c, area[(r, c)]) for c in cols if (r, c) in area]
            toks = [(c, v) for c, v in cells if X.MTOK.match(X.txt(v)) and X.to_num(v) is None]
            lab = X.txt(cells[0][1]) if cells else ''
            labok = (not cells) or lab in ('-', '–') or X.MTOK.match(lab) or bool(X.NUMLABEL.match(lab.rstrip(':*'))) or lab.lower().startswith(('thread', 'tdread', 'závit', 'size', 'for thread', 'nominal'))
            if (len(toks) >= 2 and labok) or (X.NUMLABEL.match(lab.rstrip(':*')) and lab[:1].isalpha() and len(cells) >= 4 and
                                   sum(X.to_num(v) is not None for c, v in cells[1:]) >= len(cells) - 2 and r == b[0]):
                size_rows.append(r)
        if size_rows:
            def lab_of(r):
                cs = [X.txt(area[(r, c)]) for c in cols if (r, c) in area]
                return cs[0] if cs else ''
            keep = [size_rows[0]]
            for r in size_rows[1:]:
                if r == keep[-1] + 1 or lab_of(r) == lab_of(size_rows[0]):
                    keep.append(r)
            size_rows = keep
        # a later size header whose sizes span two columns with text sub-labels below (ls / lg) starts a grouped grid
        if size_rows:
            for r in size_rows:
                spans = [sh.merge_of.get((r, c)) for c in cols if (r, c) in area and X.MTOK.match(X.txt(area[(r, c)]))]
                wide = sum(1 for m in spans if m and m[3] > m[1])
                nxt = [X.txt(area[(r + 1, c)]) for c in cols if (r + 1, c) in area]
                texty = sum(1 for t in nxt[1:] if t and not isnum(clean(t)) and not X.MTOK.match(t.strip('()')) and t not in ('-', '–'))
                scs = sorted(c for c in cols if (r, c) in area and X.MTOK.match(X.txt(area[(r, c)])))
                gapped = sum(1 for x, y in zip(scs, scs[1:]) if y - x == 2)
                nxt2 = [X.txt(area[(r + d, c)]) for d in (2, 3, 4) for c in cols if (r + d, c) in area]
                lsg = any(re.match(r'^(ls|lg)\b', t) for t in nxt + nxt2) or any('shank' in t.lower() for t in nxt + nxt2)
                if (wide >= 2 and texty >= max(2, len(nxt) // 2)) or ((wide >= 2 or gapped >= 2) and lsg):
                    k = b.index(r)
                    g_rows = b[k:]
                    b = b[:k]; size_rows = [x for x in size_rows if x < r]
                    extra_g = g_table(sh, area, g_rows, sorted({c for (rr, c) in area if rr in g_rows}), isint, fixes)
                    break
            else:
                extra_g = None
        else:
            extra_g = None
        if not b:
            if extra_g is not None: tables.append(extra_g)
            continue
        if size_rows and size_rows[0] - b[0] <= 1:
            # sizes must own the numbers: many numbers between size columns means a grouped grid (e.g. ls/lg tables)
            sr = size_rows[0]
            hdr0 = [(c, area[(sr, c)]) for c in cols if (sr, c) in area]
            sc0 = {c for c, v in hdr0 if X.MTOK.match(X.txt(v)) or X.to_num(v) is not None}
            for c, v in hdr0:
                m = sh.merge_of.get((sr, c))
                if m and c in sc0: sc0 |= set(range(c, m[3] + 1))
            first_sc = min(sc0) if sc0 else 0
            inside = outside = 0
            for r in b:
                if r <= sr or r in size_rows: continue
                for c in cols:
                    if (r, c) in area and isnum(clean(area[(r, c)])):
                        if c in sc0: inside += 1
                        elif c > first_sc: outside += 1
            if outside > 0.2 * max(1, inside + outside):
                size_rows = []
        if size_rows and size_rows[0] - b[0] <= 1:
            # a trailing grid (e.g. length tolerances) starts where the label column carries numbers
            hdr = [(c, area[(size_rows[0], c)]) for c in cols if (size_rows[0], c) in area]
            sc = [c for c, v in hdr if X.MTOK.match(X.txt(v)) or X.to_num(v) is not None]
            labc = [c for c in cols if sc and c < sc[0]]
            cut = None
            for r in b:
                if r <= size_rows[0] or r in size_rows:
                    continue
                if labc and (r, labc[0]) in area and isnum(clean(area[(r, labc[0])])):
                    k = b.index(r)
                    while k - 1 >= 0 and b[k - 1] not in size_rows and not any(isnum(clean(area[(b[k - 1], c)])) for c in cols if (b[k - 1], c) in area):
                        k -= 1
                    cut = k; break
            hb = b[:cut] if cut is not None else b
            tables += h_tables(sh, area, hb, cols, [r for r in size_rows if r in hb], isint, fixes)
            if cut is not None:
                gb = b[cut:]
                tables.append(g_table(sh, area, gb, sorted({c for (r, c) in area if r in gb}), isint, fixes))
        else:
            tables.append(g_table(sh, area, b, cols, isint, fixes))
        if b and b[0] in caps and len(tables) > n0:
            tables[n0]['title'] = caps[b[0]]
        if extra_g is not None:
            tables.append(extra_g)
    # join consecutive h tables that share row structure (wrapped / non-standard size blocks)
    out = []
    for t in tables:
        prevh = next((x for x in reversed(out) if x['k'] == 'h'), None)
        if t['k'] == 'g' and any(x['k'] == 'g' and x['head'] == t['head'] and x['body'] == t['body'] for x in out):
            DUP[0] += sum(isnum(x) for line in t['body'] for x in line)
            continue  # the source repeats the same grid under a second size block
        if t['k'] == 'h' and prevh is not None and not t.get('title') and not (set(t['sizes']) & set(prevh['sizes'])):
            p = prevh
            if [(x['sym'], x['q']) for x in p['rows']] == [(x['sym'], x['q']) for x in t['rows']] or \
               len({(x['sym'], x['q']) for x in t['rows']} & {(x['sym'], x['q']) for x in p['rows']}) >= 0.6 * len(t['rows']):
                keys = [(x['sym'], x['q']) for x in p['rows']]
                for x in t['rows']:
                    if (x['sym'], x['q']) not in keys:
                        p['rows'].append({'sym': x['sym'], 'q': x['q'], 'v': [''] * len(p['sizes'])}); keys.append((x['sym'], x['q']))
                for i, z in enumerate(t['sizes']):
                    p['sizes'].append(z); p['ns'].append(t['ns'][i])
                    for row in p['rows']:
                        src = next((x for x in t['rows'] if (x['sym'], x['q']) == (row['sym'], row['q'])), None)
                        row['v'].append(src['v'][i] if src else '')
                continue
        out.append(t)
    return out, fixes

def h_tables(sh, area, b, cols, size_rows, isint, fixes):
    res = []
    seg = [r for i, r in enumerate(size_rows) if i == 0 or size_rows[i - 1] != r - 1]
    alias = set(size_rows) - set(seg)
    for si, R in enumerate(seg):
        end = seg[si + 1] if si + 1 < len(seg) else b[-1] + 1
        hdr = [(c, area[(R, c)]) for c in cols if (R, c) in area]
        scols = [(c, v) for c, v in hdr if X.MTOK.match(X.txt(v)) or X.to_num(v) is not None]
        if len(scols) < 2:
            continue
        # a size header merged over two columns carries a second variant (e.g. ISO / DIN wrench size)
        var = {}
        for c, v in list(scols):
            m = sh.merge_of.get((R, c))
            if m and m[3] > c:
                for c2 in range(c + 1, m[3] + 1):
                    scols.append((c2, v)); var[c2] = c2 - c
        scols.sort()
        lab_cols = [c for c in cols if c < scols[0][0]]
        nonstd = any('nonstandard' in X.txt(v).lower() for c, v in hdr)
        sizes = [X.norm_size(v)[0] + ('′″‴'[min(var[c], 3) - 1] if c in var else '') for c, v in scols]
        rows, prev_sym = [], ''
        for r in range(R + 1, end):
            if r not in b:
                continue
            labs = [X.txt(area[(r, c)]) for c in lab_cols if (r, c) in area]
            vals = []
            for c, v in scols:
                cell = area.get((r, c))
                if cell is None and (r, c) in sh.merge_of:  # merged value spans several sizes
                    m = sh.merge_of[(r, c)]; cell = sh.raw.get((m[0], m[1])) if m[0] == r else None
                vals.append(clean(cell) if cell is not None else '')
            if not any(vals) and not labs:
                continue
            if r in alias:
                rows.append({'sym': 'd', 'q': 'رزوهٔ ریزدنده', 'v': [x.replace(',', '.').replace('x', '×').replace('X', '×') if x not in ('-', '') else x for x in vals], '_r': r})
                continue
            if not labs:
                sym, q = prev_sym, ''
            elif len(lab_cols) >= 2 and (r, lab_cols[0]) not in area and prev_sym and not any((r, c) in area for c in lab_cols[:-1] if c != lab_cols[-1] and c == lab_cols[0]):
                sym, q = prev_sym, ' '.join(labs)
            else:
                first = labs[0]
                if first.lower().startswith(('min', 'max', 'nominal', 'nom.', 'production', 'approx', 'per', '≈', 'tolerance', 'mix')):
                    sym, q = prev_sym, ' '.join(labs)
                else:
                    sym, q = first, ' '.join(labs[1:])
            prev_sym = sym
            rows.append({'sym': sym, 'q': q, 'v': vals, '_r': r})
        # lost decimal comma: compare with sibling rows of the same symbol in the same size column
        for i, row in enumerate(rows):
            for j, t in enumerate(row['v']):
                if not isnum(t) or (row['_r'], scols[j][0]) not in isint:
                    continue
                v = num(t)
                if v is None or v < 100:
                    continue
                sib = [num(x['v'][j]) for x in rows if x['sym'] == row['sym'] and x is not row and isnum(x['v'][j])]
                sib = [s for s in sib if s is not None and s < v / 20]
                if not sib:
                    # single-row symbol: compare with the whole size column (weights excluded), only for gross outliers
                    colv = [num(x['v'][j]) for x in rows if x is not row and isnum(x['v'][j]) and not re.search(r'weight|kg|kn', x['sym'] + ' ' + x['q'], re.I)]
                    colv = [s for s in colv if s is not None]
                    if len(colv) < 3 or v / statistics.median(colv) < 50 or re.search(r'weight|kg|kn', row['sym'] + ' ' + row['q'], re.I):
                        continue
                    sib = colv
                m = statistics.median(sib)
                best = min((abs(v / 10 ** k - m), k) for k in (1, 2, 3, 4))
                cand = v / 10 ** best[1]
                if 0.5 * m <= cand <= 2 * m:
                    row['v'][j] = ('%g' % cand); fixes.append('%s %s %s: %s→%s' % (row['sym'], row['q'], sizes[j], t, row['v'][j]))
        # impossible pitch values
        for row in rows:
            if row['sym'] == 'P':
                for j, t in enumerate(row['v']):
                    z = re.match(r'^M(\d+(?:\.\d+)?)$', sizes[j])
                    if z and isnum(t) and z.group(1) in PITCH and num(t) not in PITCH[z.group(1)]:
                        new = '%g' % PITCH[z.group(1)][0]; fixes.append('P %s: %s→%s (ISO 261)' % (sizes[j], t, new)); row['v'][j] = new
        for row in rows:
            row.pop('_r', None)
        res.append({'k': 'h', 'sizes': sizes, 'ns': [nonstd] * len(sizes), 'rows': rows})
    return res

def g_table(sh, area, b, cols, isint, fixes):
    # header rows: leading rows until the first row that is mostly numeric
    def numeric_ratio(r):
        cs = [clean(area[(r, c)]) for c in cols if (r, c) in area]
        return sum(isnum(x) for x in cs) / max(1, len(cs)), len(cs)
    hn = 0
    for r in b:
        ratio, n = numeric_ratio(r)
        if ratio >= 0.5 and n >= 2:
            break
        hn += 1
    hn = min(hn, 6)
    head_rows, body_rows = b[:hn], b[hn:]
    c0, c1 = cols[0], cols[-1]
    head = []
    for r in head_rows:
        line, c = [], c0
        while c <= c1:
            m = sh.merge_of.get((r, c))
            if m and (m[0], m[1]) != (r, c):
                c += 1; continue
            t = X.txt(area.get((r, c), ''))
            cs = (min(m[3], c1) - c + 1) if m else 1
            rs = (min(m[2], head_rows[-1]) - r + 1) if m and head_rows else 1
            line.append({'t': t, 'fa': qfa(t), 'cs': cs, 'rs': max(1, rs)})
            c += cs
        head.append(line)
    body = []
    for r in body_rows:
        line = []
        for c in range(c0, c1 + 1):
            v = area.get((r, c))
            if v is None and (r, c) in sh.merge_of:
                m = sh.merge_of[(r, c)]
                if m[0] == r:  # horizontal merge -> repeat
                    v = sh.raw.get((m[0], m[1]))
            line.append(clean(v) if v is not None else '')
        if any(line):
            body.append(line)
    # lost decimal comma inside a column: compare with neighbouring rows
    for j in range(c1 - c0 + 1):
        colv = [num(x[j]) if isnum(x[j]) else None for x in body]
        for i, v in enumerate(colv):
            if v is None or v < 100 or (body_rows[min(i, len(body_rows) - 1)], c0 + j) not in isint:
                continue
            nb = [colv[k] for k in range(max(0, i - 3), min(len(colv), i + 4)) if k != i and colv[k] is not None]
            if len(nb) < 2:
                continue
            m = statistics.median(nb)
            if m < v / 20:
                best = min((abs(v / 10 ** k - m), k) for k in (1, 2, 3, 4)); cand = v / 10 ** best[1]
                if 0.5 * m <= cand <= 2 * m:
                    body[i][j] = '%g' % cand; colv[i] = cand; fixes.append('col%d row%d: %g→%g' % (j + 1, i + 1, v, cand))
    return {'k': 'g', 'head': head, 'body': body}

def audit(sh, area_tables, area):
    """numeric cells of the source table area vs numbers present in the output"""
    src = collections.Counter()
    for k, v in area.items():
        t = clean(v)
        if isnum(t):
            src[t.strip('()+')] += 1
    got = collections.Counter()
    for t in area_tables:
        if t['k'] == 'h':
            for z in t['sizes']:
                got[z.lstrip('M').replace('M', '')] += 1
            for row in t['rows']:
                for x in row['v']:
                    if isnum(x): got[x.strip('()+')] += 1
        else:
            for line in t['body']:
                for x in line:
                    if isnum(x): got[x.strip('()+')] += 1
            for line in t['head']:
                for x in line:
                    if isnum(clean(x['t'])): got[clean(x['t']).strip('()+')] += 1
    miss = sum((src - got).values())
    return sum(src.values()), miss

def main():
    parents = load_parents()
    out, report = {}, []
    for f in sorted(glob.glob(ROOT + '/data/raw/*.xlsx')):
        w = openpyxl.load_workbook(f, data_only=True)
        for ws in w.worksheets[1:]:
            if ws.title.strip().lower() == 'sheet1':
                continue
            sh = X.Sheet(ws, f)
            if not sh.code:
                continue
            code = (sh.code + (sh.form or '')).replace(' ', '').upper()
            pc = next((k for k in parents if k.upper() == code or parents[k]['Source sheet'] == ws.title), None)
            if not pc or pc in out:
                continue
            DUP[0] = 0
            tables, fixes = sheet_tables(sh, ws)
            area = {k: v for k, v in sh.raw.items() if k not in sh.info_cells and k[1] < sh.info_col}
            n, miss = audit(sh, tables, area)
            legend = {}
            for t in sh.legend:
                m = re.match(r"^(\S+(?:\s\S+)?)\s+-\s+(.+)$", t)
                if m: legend[m.group(1)] = m.group(2)
            out[pc] = {'tables': tables, 'legend': legend, 'fixes': fixes, 'sheet': ws.title}
            report.append((pc, n, max(0, miss - len(fixes) - DUP[0]), len(tables), sum(len(t['rows']) for t in tables if t['k'] == 'h'), len(fixes)))
    json.dump(out, open(OUTDIR + '/DIN_full_tables.json', 'w'), ensure_ascii=False, indent=0)
    tot = sum(r[1] for r in report); miss = sum(r[2] for r in report)
    print('sheets', len(report), 'numeric cells', tot, 'missing', miss, 'fixes', sum(r[5] for r in report))
    for r in sorted(report, key=lambda r: -r[2])[:25]:
        if r[2]: print('MISS', r)
    json.dump(report, open(OUTDIR + '/DIN_full_tables_audit.json', 'w'))

if __name__ == '__main__':
    main()
