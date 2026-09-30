"""Extract unique size-level SKUs from the DIN detail workbooks.

Each worksheet in data/raw/*.xlsx holds the dimension tables of one DIN
standard (scraped from fasteners.eu). Layouts differ from sheet to sheet, so
the parser detects the size axis (horizontal header row or vertical column),
then looks for a length table (nominal length rows x size columns).

One SKU = standard (+ form) x size [x length]. Material, property class and
coating are not in the source data and are left for a later layer.

Usage: python3 scripts/extract_skus.py  ->  output/DIN_SKU_list.xlsx
"""

import datetime
import glob
import math
import os
import re
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
OUT = os.path.join(ROOT, "output", "DIN_SKU_list.xlsx")

# Thread-type size tokens: M8, M8x1, (M14), ST3,5, Tr10x2, G1/4, R1/2, 1/4"
MTOK = re.compile(
    r"^\(?\s*((?:MB)\s*\d+"
    r"|(?:M|m|ST|Tr|TR)\s*\d+(?:[,.]\d+)?(?:\s*[x×]\s*\d+(?:[,.]\d+)?)?"
    r"|(?:G|R|Rp)\s*\d+(?:\s*\d+/\d+|/\d+)?"
    r"|\d+(?:[,.]\d+)?\s*[x×]\s*\d+(?:[,.]\d+)?"
    r"|\d+/\d+\s*(?:\"|'')?)\s*\)?$"
)
NUM = re.compile(r"^\(?\s*-?\d+(?:[,.]\d+)?\s*\)?$")
# Labels of a row/column that carries the nominal size when sizes are plain numbers
NUMLABEL = re.compile(
    r"^(?:d\d?|d\s?1|ds|d\*|De|D|(?i:d\s*nom.*|size.*|nominal.*|dimension|"
    r"shaft diameter|bore diameter|diameter.*|thread.*|závit.*|for thread.*|"
    r"pin diameter|rivet diameter))(?:\s+[a-zA-Z]{1,2}\d{1,2})?$"
)
# last-resort size labels (keys: b, hex keys: s)
EXTLABEL = re.compile(r"^(?:b|s|b x h)(?:\s+[a-zA-Z]{1,2}\d{1,2})?$")
SECONDARY = re.compile(r"^d[2-9]\b")
RANGE = re.compile(r"^\d+(?:[,.]\d+)?\s*-\s*\d+(?:[,.]\d+)?$")
# Products whose length is fixed by the size (eyebolts, keys, grease nipples):
# the legend mentions a length but it is not a variant axis.
FIXED_LENGTH = {"DIN 580", "DIN 6888", "DIN 71412 a"}
TITLE = re.compile(r"^(DIN)\s*(\d+)\s*([a-z]{0,4})\s*-\s*(.+)$", re.I)


def txt(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):  # Excel turned "8-12" into a date
        return "%d-%d" % (v.month, v.day)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def to_num(v):
    """Parse '2,5', 2.5, '(35)', -35 -> float. Returns None if not numeric."""
    if isinstance(v, (int, float)):
        return float(v)
    s = txt(v)
    if not NUM.match(s):
        return None
    return float(s.strip("() ").replace(",", "."))


def is_size_num(v):
    return to_num(v) is not None or bool(RANGE.match(txt(v)))


def fix_scale(values):
    """Repair numbers whose decimal comma was lost (2,5 stored as 25).
    values: list of floats in table order. Returns (fixed, notes)."""
    out, notes = list(values), []
    for i, v in enumerate(out):
        prev = out[i - 1] if i > 0 else None
        nxt = out[i + 1] if i + 1 < len(out) else None
        if prev is None or v <= prev * 3:
            continue
        for k in (10, 100, 1000):
            cand = v / k
            if cand >= prev and (nxt is None or cand <= nxt):
                if nxt is not None or cand <= prev * 3:
                    out[i] = cand
                    notes.append("%s→%s" % (fmt_num(v), fmt_num(cand)))
                    break
    return out, notes


def nominal_d(size):
    m = re.match(r"^(?:M|ST|TR)(\d+(?:\.\d+)?)", size)
    return float(m.group(1)) if m else None


def fmt_num(x):
    return str(int(x)) if float(x).is_integer() else ("%g" % x)


def norm_size(tok):
    """'(M1,6)' -> ('M1.6', nonpreferred=True); 2.5 -> ('2.5', False)."""
    s = txt(tok)
    nonpref = s.startswith("(") and s.endswith(")")
    s = s.strip("() ")
    n = to_num(s)
    if n is not None:
        return fmt_num(abs(n)), nonpref or n < 0
    s = s.replace(",", ".").replace("×", "x").replace(" ", "")
    s = re.sub(r"^(m|st|tr|g|r|rp|npt)", lambda m: m.group(1).upper(), s, flags=re.I)
    s = s.replace("x", "X")
    return s, nonpref


class Sheet:
    def __init__(self, ws, source):
        self.ws = ws
        self.source = source
        self.nrow = ws.max_row
        self.ncol = ws.max_column
        self.raw = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and txt(c.value) != "":
                    self.raw[(c.row, c.column)] = c.value
        # thread written over two cells: "M" above "10x1"
        for (r, c), v in list(self.raw.items()):
            below = txt(self.raw.get((r + 1, c)))
            if txt(v) == "M" and re.match(
                r"^\d+(?:[,.]\d+)?(?:x\d+(?:[,.]\d+)?)?$", below
            ):
                self.raw[(r, c)] = "M" + below
                del self.raw[(r + 1, c)]
        # merged ranges: every covered cell points to its range
        self.merge_of = {}
        for mr in ws.merged_cells.ranges:
            rng = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    self.merge_of[(r, c)] = rng
        self.warnings = []
        self.find_meta()

    # value with merged cells filled from their top-left cell
    def val(self, r, c):
        if (r, c) in self.raw:
            return self.raw[(r, c)]
        rng = self.merge_of.get((r, c))
        if rng:
            return self.raw.get((rng[0], rng[1]))
        return None

    def row_cells(self, r, maxc=None):
        maxc = maxc or self.ncol
        return [
            (c, self.raw[(r, c)])
            for c in range(1, maxc + 1)
            if (r, c) in self.raw and (r, c) not in self.info_cells
        ]

    def find_meta(self):
        self.title = self.code = self.name = self.form = None
        self.current_norm = self.equivalents = ""
        self.info_col = self.ncol + 1
        self.legend = []
        for (r, c), v in sorted(self.raw.items()):
            s = txt(v)
            m = TITLE.match(s)
            if m and self.title is None and " - " in s:
                self.title = s
                self.code = "DIN %s" % m.group(2)
                self.form = m.group(3).lower()
                self.name = m.group(4).strip()
                self.info_col = c
            elif s.startswith("Current norm"):
                self.current_norm = s.split(":", 1)[1].strip()
            elif s.startswith("Equivalent norms"):
                self.equivalents = s.split(":", 1)[1].strip().rstrip(";")
        if self.title is None:
            m = re.match(r"^DIN\s*(\d+)\s*([a-z]*)$", self.ws.title.strip(), re.I)
            if m:
                self.code = "DIN %s" % m.group(1)
                self.form = m.group(2).lower()
                self.name = ""
                self.warnings.append(
                    "عنوان استاندارد در شیت نیست؛ کد از نام شیت گرفته شد"
                )
        # legend lines ("l - length of bolt") live in the info column
        self.info_cells = set()
        for (r, c), v in self.raw.items():
            t = txt(v)
            if c >= self.info_col and re.match(r"^\S+(\s\S+)?\s+-\s+\D", t):
                self.legend.append(t)
            if (
                (TITLE.match(t) and " - " in t)
                or t.startswith(("Current norm", "Equivalent norms"))
                or (c >= self.info_col and re.match(r"^\S+(\s\S+)?\s+-\s+\D", t))
                or "dimensions in mm" in t
                or t.startswith("Table according")
            ):
                self.info_cells.add((r, c))
        self.has_length_attr = any(
            re.match(r"^[lL]\d?\s+-\s+.*length", x) for x in self.legend
        )

    # ---------- size axis detection ----------
    def h_headers(self):
        """Rows with >=2 thread-type tokens: [(row, label_col, [(col, token)])]."""
        out = []
        for r in range(1, self.nrow + 1):
            cells = self.row_cells(r)
            toks = [(c, v) for c, v in cells if MTOK.match(txt(v))]
            if len(toks) >= 2:
                lab = [c for c, v in cells if c < toks[0][0]]
                out.append((r, lab[0] if lab else toks[0][0], toks))
        return out

    def v_column(self):
        """Column with the most thread-type tokens stacked vertically."""
        best = (0, None)
        for c in range(1, self.ncol + 1):
            n = sum(
                1
                for r in range(1, self.nrow + 1)
                if (r, c) in self.raw and MTOK.match(txt(self.raw[(r, c)]))
            )
            if n > best[0]:
                best = (n, c)
        return best

    def h_numeric_header(self, labels=NUMLABEL):
        """First row labelled like a size (d, d1, Size, ...) followed by numbers.
        Later rows with the same label continue the size list (wrapped tables)."""
        out, label = [], None
        for r in range(1, self.nrow + 1):
            cells = self.row_cells(r)
            if not cells:
                continue
            c0, v0 = cells[0]
            lab = txt(v0).rstrip(":*")
            if to_num(v0) is not None or not labels.match(lab):
                continue
            if label is not None and lab != label:
                continue
            rest = cells[1:]
            first_num = next(
                (i for i, (c, v) in enumerate(rest) if to_num(v) is not None), None
            )
            if first_num is None:
                continue
            nums = [(c, v) for c, v in rest[first_num:] if to_num(v) is not None]
            # sub-labels ("before mounting", "nominal") may precede the numbers only
            if len(nums) >= 2 and len(nums) == len(rest) - first_num and first_num <= 2:
                out.append((r, c0, nums))
                label = lab
        return out

    def v_numeric_column(self, labels=NUMLABEL):
        for r in range(1, self.nrow + 1):
            cells = self.row_cells(r)
            if len(cells) < 3:
                continue
            for c, v in cells:
                if to_num(v) is None and labels.match(txt(v).rstrip(":*")):
                    below = [
                        self.raw.get((rr, c))
                        for rr in range(r + 1, min(r + 6, self.nrow + 1))
                    ]
                    if sum(is_size_num(x) for x in below) >= 2:
                        return r, c
                    break
        return None

    def numeric_row_above(self, hr, cols):
        """Plain numeric size row above a secondary thread row. Only rows that are
        unlabelled or labelled like a nominal size (d, d1, ...) qualify."""
        for r in range(hr - 1, max(hr - 4, 0), -1):
            nums = [
                (c, self.raw.get((r, c)))
                for c in cols
                if to_num(self.raw.get((r, c))) is not None
            ]
            if len(nums) < max(2, len(cols) - 1):
                continue
            labs = [(c, v) for c, v in self.row_cells(r) if c < cols[0]]
            lab = txt(labs[0][1]) if labs else ""
            if lab and (not NUMLABEL.match(lab) or SECONDARY.match(lab)):
                continue
            return (r, labs[0][0] if labs else cols[0], nums)
        return None

    def numeric_column_left(self, vcol):
        """Numeric 'd'/'d1' column left of a thread column (ring bore, knob size)."""
        first = next(
            (
                r
                for r in range(1, self.nrow + 1)
                if (r, vcol) in self.raw and MTOK.match(txt(self.raw[(r, vcol)]))
            ),
            None,
        )
        if not first:
            return None
        # washers are sold by the bolt they fit ("For thread" M8), keep that axis
        if any(
            "for thread" in txt(self.raw.get((r, vcol))).lower()
            for r in range(max(first - 3, 1), first)
        ):
            return None
        for hdr in range(first - 1, max(first - 4, 0), -1):
            for c in range(1, vcol):
                h = txt(self.raw.get((hdr, c)))
                if re.match(r"^d1?(\s|$)", h):
                    below = [
                        self.raw.get((r, c))
                        for r in range(first, min(first + 5, self.nrow + 1))
                    ]
                    rng = self.merge_of.get((hdr, c))
                    if rng and rng[3] > c:
                        below += [
                            self.raw.get((r, c + 1))
                            for r in range(first, min(first + 5, self.nrow + 1))
                        ]
                    if sum(to_num(x) is not None for x in below) >= 3:
                        return hdr, c
        return None

    # ---------- extraction ----------
    def spans(self, header_row, toks):
        spans = []
        for i, (c, tok) in enumerate(toks):
            rng = self.merge_of.get((header_row, c))
            if rng:
                end = rng[3]
            elif i + 1 < len(toks):
                end = toks[i + 1][0] - 1
            else:
                end = c + (spans[-1][2] - spans[-1][1] if spans else 0)
            end = min(end, self.ncol)
            spans.append((tok, c, end))
        return spans

    def pitch_row(self, header_row, spans):
        """Pitch per size from a row labelled 'P' right under a header."""
        out = {}
        for r in range(header_row + 1, min(header_row + 6, self.nrow + 1)):
            cells = self.row_cells(r)
            if cells and txt(cells[0][1]) in ("P", "P*", "Pitch", "P (pitch)"):
                for tok, c0, c1 in spans:
                    v = to_num(self.val(r, c0))
                    if v is not None:
                        out[norm_size(tok)[0]] = fmt_num(v)
                return out
        return out

    def extract(self):
        """Returns (method, sizes OrderedDict, combos list, length_list)."""
        sizes = OrderedDict()  # size -> dict(nonpref, pitch)
        combos = []  # (size, length, nonpref_len, confidence)
        lengths_seen = []
        hh = self.h_headers()
        vcount, vcol = self.v_column()
        hsizes = max((len(t) for _, _, t in hh), default=0)
        headers, method = [], None
        if hh and hsizes >= vcount:
            method, headers = "H", hh
            # a secondary thread row (d2/d3 = internal thread of a pin) is not the
            # size axis when a plain numeric size row sits right above it
            hr, lc, toks = hh[0]
            if SECONDARY.match(txt(self.raw.get((hr, lc)))):
                above = self.numeric_row_above(hr, [c for c, _ in toks])
                if above:
                    method, headers = "H-num", [above]
        elif vcount >= 3:
            left = self.numeric_column_left(vcol)
            if left:
                return self.extract_vertical(left[1], "V-num", header_row=left[0])
            return self.extract_vertical(vcol, "V")
        if not headers:
            for labels in (NUMLABEL, EXTLABEL):
                headers = self.h_numeric_header(labels)
                if headers:
                    method = "H-num"
                    break
                vn = self.v_numeric_column(labels)
                if vn:
                    return self.extract_vertical(vn[1], "V-num", header_row=vn[0])
            if not headers:
                return "none", sizes, combos, lengths_seen
        if method == "H-num":
            flat = [abs(to_num(v)) for _, _, toks in headers for _, v in toks]
            fixed, notes = fix_scale(flat)
            if notes:
                self.warnings.append("اصلاح اعشار گم‌شده در سایز: " + ", ".join(notes))
                it = iter(fixed)
                headers = [
                    (hr, lc, [(c, math.copysign(next(it), to_num(v))) for c, v in toks])
                    for hr, lc, toks in headers
                ]

        # stacked header rows (coarse / fine / extra-fine thread) share one column
        groups = []
        for h in headers:
            if groups and h[0] == groups[-1][-1][0] + 1:
                groups[-1].append(h)
            else:
                groups.append([h])
        for gi, group in enumerate(groups):
            hr0 = group[0][0]
            base = max(group, key=lambda h: len(h[2]))
            spans = [[a, b, []] for _, a, b in self.spans(base[0], base[2])]
            for hr, labc, toks in group:
                nonstd_block = any(
                    "nonstandard" in txt(v).lower() for _, v in self.row_cells(hr)
                )
                own = self.spans(hr, toks)
                pitches = self.pitch_row(group[-1][0], own) if hr == base[0] else {}
                for tok, a, b in own:
                    s_, npf = norm_size(tok)
                    if s_ not in sizes:
                        sizes[s_] = {
                            "nonpref": npf or nonstd_block,
                            "pitch": pitches.get(s_, ""),
                        }
                    elif pitches.get(s_) and not sizes[s_]["pitch"]:
                        sizes[s_]["pitch"] = pitches[s_]
                    home = next((sp for sp in spans if sp[0] <= a <= sp[1]), None)
                    if home is None:
                        home = [a, b, []]
                        spans.append(home)
                    home[2].append(s_)
            first_size_col = min(sp[0] for sp in spans)
            next_hr = groups[gi + 1][0][0] if gi + 1 < len(groups) else self.nrow + 1
            group_combos, group_lengths = [], []
            for r in range(group[-1][0] + 1, next_hr):
                cells = self.row_cells(r)
                if not cells:
                    continue
                c0, v0 = cells[0]
                L = to_num(v0)
                if L is None or c0 >= first_size_col or L == 0:
                    continue
                Lnonpref = (isinstance(v0, (int, float)) and v0 < 0) or txt(
                    v0
                ).startswith("(")
                L = abs(L)
                if L not in [x[0] for x in group_lengths]:
                    group_lengths.append((L, Lnonpref))
                for a, b, toks in spans:
                    vals, explicit = [], False
                    for c in range(a, b + 1):
                        v = self.val(r, c)
                        if v is None or txt(v) == "" or (r, c) in self.info_cells:
                            continue
                        vals.append(txt(v))
                        rng = self.merge_of.get((r, c))
                        # a merged cell that spills over other sizes is ambiguous
                        if rng is None or (rng[1] >= a and rng[3] <= b):
                            explicit = True
                    if not vals or all(x in ("-", "–") for x in vals):
                        continue
                    for s_ in toks:
                        if not explicit:
                            d = nominal_d(s_)
                            # physical sanity check: a screw shorter than 1.5 x d is not
                            # offered (matches ISO 4762 / 4017 minimum lengths)
                            if d and L < 1.5 * d:
                                continue
                        group_combos.append(
                            (s_, L, Lnonpref, "explicit" if explicit else "ambiguous")
                        )
            # a real length table has several rows; a stray number is not one
            if len(group_lengths) >= 3:
                combos.extend(group_combos)
                for x in group_lengths:
                    if x[0] not in [y[0] for y in lengths_seen]:
                        lengths_seen.append(x)
        return method, sizes, combos, lengths_seen

    def extract_vertical(self, col, method, header_row=None):
        sizes = OrderedDict()
        combos, lengths_seen = [], []
        start = header_row + 1 if header_row else 1
        # a column literally headed "l" next to the size column = length per row
        lcol, pcol = None, None
        hdr = header_row
        if hdr is None:
            for r in range(1, self.nrow + 1):
                if (r, col) in self.raw and MTOK.match(txt(self.raw[(r, col)])):
                    hdr = r - 1
                    break
        if hdr:
            for c in range(1, self.ncol + 1):
                h = txt(self.val(hdr, c)).lower()
                if c != col and re.match(r"^l1?\s*(nom.*|nominal)?$", h):
                    lcol = c
                if c != col and h in ("p", "pitch", "p*"):
                    pcol = c
        # header merged over two columns = "Series 1 | Series 2" size columns
        rng = self.merge_of.get((hdr, col)) if hdr else None
        alt_col = col + 1 if rng and rng[3] > col else None
        rows, blank = [], 0
        for r in range(max(start, (hdr or 0) + 1), self.nrow + 1):
            if not self.row_cells(r):
                blank += 1
                if blank >= 2 and rows:
                    break
                continue
            blank = 0
            v = self.raw.get((r, col))
            if v is None and alt_col:
                v = self.raw.get((r, alt_col))
            if v is None:
                rows.append((r, None))
                continue
            ok = MTOK.match(txt(v)) if method == "V" else is_size_num(v)
            if ok:
                rows.append((r, v))
        if method == "V-num":
            nums = [
                (i, abs(to_num(v)))
                for i, (r, v) in enumerate(rows)
                if v is not None and to_num(v) is not None
            ]
            fixed, notes = fix_scale([x for _, x in nums])
            if notes:
                self.warnings.append("اصلاح اعشار گم‌شده در سایز: " + ", ".join(notes))
                for (i, _), fx in zip(nums, fixed):
                    rows[i] = (rows[i][0], fx)
        keys = [norm_size(v)[0] for r, v in rows if v is not None]
        dup = len(keys) != len(set(keys)) and not lcol
        if dup:
            self.warnings.append("سایز تکراری در ستون اصلی؛ کلید با ستون بعدی ترکیب شد")
        cur = None
        for r, v in rows:
            if v is not None:
                s, npf = norm_size(v)
                if dup and keys.count(s) > 1:
                    for extra in range(col + 1, col + 4):
                        e = self.val(r, extra)
                        if e is not None and txt(e) not in ("", "-"):
                            s += "×" + norm_size(e)[0]
                            if s not in sizes:
                                break
                cur = s
                if s not in sizes:
                    pitch = to_num(self.val(r, pcol)) if pcol else None
                    sizes[s] = {
                        "nonpref": npf,
                        "pitch": fmt_num(pitch) if pitch else "",
                    }
            if lcol and cur:
                L = to_num(self.val(r, lcol))
                if L:
                    combos.append((cur, abs(L), False, "explicit"))
                    if abs(L) not in [x[0] for x in lengths_seen]:
                        lengths_seen.append((abs(L), False))
        return method, sizes, combos, lengths_seen


def sku_code(code, form, size, length=None):
    base = code.replace(" ", "") + (form.upper() if form else "")
    s = "%s-%s" % (base, size.replace("×", "_"))
    if length is not None:
        s += "X%s" % fmt_num(length)
    return s


def sku_title(std_key, size, length=None):
    t = "%s %s" % (std_key.replace(" ", " ", 1), size.replace("X", "x"))
    if length is not None:
        t += " × %s" % fmt_num(length)
    return t


def confidence(method, warnings, status):
    if method == "none":
        return "—"
    if warnings or status.startswith("طول‌ها در منبع"):
        return "پایین"
    if method in ("H-num", "V-num"):
        return "متوسط"
    return "بالا"


def main():
    files = sorted(glob.glob(os.path.join(RAW, "*.xlsx")))
    skus, review, standards = [], [], []
    seen_std = {}
    for f in files:
        part = re.search(r"part(\d)", os.path.basename(f)).group(1)
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets[1:]:
            if ws.title.strip().lower() == "sheet1":
                continue  # priority index sheet, not a standard
            sh = Sheet(ws, f)
            if not sh.code:
                continue
            method, sizes, combos, lengths = sh.extract()
            std_key = sh.code + (" " + sh.form if sh.form else "")
            if std_key in FIXED_LENGTH:
                sh.has_length_attr = False
            if std_key in seen_std:
                sh.warnings.append("تکراری؛ قبلاً در Part %s آمده" % seen_std[std_key])
            seen_std.setdefault(std_key, part)

            base = dict(
                standard=std_key,
                base_standard=sh.code,
                form=sh.form or "",
                name_en=sh.name or "",
                current_norm=sh.current_norm,
                equivalents=sh.equivalents,
                priority="Part %s" % part,
                sheet=ws.title,
            )
            explicit = [c for c in combos if c[3] == "explicit"]
            ambiguous = [c for c in combos if c[3] == "ambiguous"]

            sizes_with_len = {c[0] for c in explicit}

            def basis_of(conf):
                if conf == "explicit":
                    return "جدول قطر×طول منبع"
                return "سلول ادغام‌شده («full thread») که روی چند قطر کشیده شده؛ مجاز بودن این ترکیب قطعی نیست"

            if combos:
                status = "قطر × طول از جدول منبع"
            elif lengths:
                status = "طول‌ها در منبع هست ولی تطبیق قطر-طول نیست"
                sh.warnings.append("جدول طول بدون نگاشت به قطر (%d طول)" % len(lengths))
            elif sh.has_length_attr:
                status = "طول جزو مشخصات است ولی در منبع نیامده"
            else:
                status = "فقط سایز (بدون طول)"

            seen = set()
            for target, group in ((skus, explicit), (review, ambiguous)):
                for s, L, Lnp, conf in sorted(
                    group,
                    key=lambda c: (
                        list(sizes).index(c[0]) if c[0] in sizes else 0,
                        c[1],
                    ),
                ):
                    code = sku_code(sh.code, sh.form, s, L)
                    if code in seen:
                        continue
                    seen.add(code)
                    target.append(
                        dict(
                            base,
                            sku=code,
                            size=s,
                            length=L,
                            pitch=sizes.get(s, {}).get("pitch", ""),
                            nonpref=(
                                "بله"
                                if (sizes.get(s, {}).get("nonpref") or Lnp)
                                else ""
                            ),
                            basis=basis_of(conf),
                        )
                    )
            # a product without a length axis is complete at size level; for
            # products that need a length, sizes without a length from the source
            # stay on the parent only (listed in the Parents sheet)
            parent_only_sizes = []
            for s, meta in sizes.items():
                if s in sizes_with_len or any(c[0] == s for c in ambiguous):
                    continue
                if status != "فقط سایز (بدون طول)":
                    parent_only_sizes.append(s)
                    continue
                code = sku_code(sh.code, sh.form, s)
                if code in seen:
                    continue
                seen.add(code)
                skus.append(
                    dict(
                        base,
                        sku=code,
                        size=s,
                        length=None,
                        pitch=meta["pitch"],
                        nonpref="بله" if meta["nonpref"] else "",
                        basis="سایز؛ این محصول طول ندارد",
                    )
                )

            if not sizes:
                sh.warnings.append("هیچ سایزی شناسایی نشد؛ بررسی دستی لازم است")
            n_children = sum(
                1
                for x in skus
                if x["standard"] == std_key and x["priority"] == base["priority"]
            )
            standards.append(
                dict(
                    base,
                    method=method,
                    n_sizes=len(sizes),
                    parent_code=sku_code(sh.code, sh.form, "").rstrip("-"),
                    role="والد + فرزند" if n_children else "فقط والد",
                    parent_only=", ".join(parent_only_sizes),
                    sizes=", ".join(sizes.keys()),
                    n_sku=sum(
                        1
                        for x in skus
                        if x["standard"] == std_key
                        and x["priority"] == base["priority"]
                    ),
                    n_review=sum(
                        1
                        for x in review
                        if x["standard"] == std_key
                        and x["priority"] == base["priority"]
                    ),
                    status=status,
                    confidence=confidence(method, sh.warnings, status),
                    lengths=", ".join(fmt_num(l) for l, _ in sorted(lengths)),
                    warnings=" | ".join(sh.warnings),
                )
            )
    conf = {(x["standard"], x["priority"]): x["confidence"] for x in standards}
    for row in skus + review:
        row["parent_code"] = sku_code(row["base_standard"], row["form"], "").rstrip("-")
        row["title"] = sku_title(row["standard"], row["size"], row["length"])
        row["confidence"] = conf.get((row["standard"], row["priority"]), "")
    write(skus, review, standards)
    print("standards:", len(standards), "skus:", len(skus), "review:", len(review))


def write(skus, review, standards):
    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")

    def sheet(title, cols, rows):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = False
        ws.append([c[1] for c in cols])
        for r in rows:
            ws.append([r.get(c[0]) if r.get(c[0]) is not None else "" for c in cols])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            ws.column_dimensions[get_column_letter(i)].width = c[2]
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    sku_cols = [
        ("sku", "SKU Code", 24),
        ("parent_code", "Parent code", 14),
        ("title", "SKU title", 26),
        ("standard", "Standard", 14),
        ("base_standard", "Base DIN", 11),
        ("form", "Form", 7),
        ("name_en", "Product name (EN)", 42),
        ("size", "Size (d)", 10),
        ("length", "Length L (mm)", 11),
        ("pitch", "Pitch P", 9),
        ("nonpref", "غیرترجیحی", 10),
        ("basis", "مبنا", 34),
        ("confidence", "اطمینان استخراج", 12),
        ("current_norm", "Current norm", 22),
        ("equivalents", "Equivalent norms", 40),
        ("priority", "Priority file", 11),
        ("sheet", "Source sheet", 12),
    ]
    std_cols = [
        ("parent_code", "Parent code", 14),
        ("standard", "Standard", 14),
        ("name_en", "Product name (EN)", 42),
        ("priority", "Priority file", 11),
        ("role", "نقش", 12),
        ("n_sizes", "# sizes", 8),
        ("n_sku", "# SKU فرزند", 9),
        ("n_review", "# برای بازبینی", 10),
        ("status", "وضعیت طول", 34),
        ("confidence", "اطمینان استخراج", 12),
        ("sizes", "Sizes", 50),
        ("parent_only", "سایزهای بدون طول (فقط در والد)", 40),
        ("lengths", "Lengths in source", 40),
        ("method", "Layout", 8),
        ("current_norm", "Current norm", 22),
        ("equivalents", "Equivalent norms", 40),
        ("sheet", "Source sheet", 12),
        ("warnings", "هشدارها", 50),
    ]
    wb.remove(wb.active)
    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    n_len = sum(1 for x in skus if x["length"] is not None)
    n_parent_only = sum(1 for x in standards if x["role"] == "فقط والد")
    lines = [
        ("محصولات DIN — والد و SKU فرزند", ""),
        ("", ""),
        (
            "مبنا",
            "فقط داده‌های DIN موجود در data/raw. از هیچ منبع بیرونی استفاده نشده است.",
        ),
        (
            "والد (Parent)",
            "هر استاندارد (و فرم آن) یک محصول والد است، مثل DIN933 یا DIN125A. برگه Parents.",
        ),
        (
            "فرزند (SKU)",
            "استاندارد × سایز [× طول]، فقط وقتی داده‌ی DIN برای ساختن آن کامل باشد. برگه SKUs. جنس، کلاس و پوشش هنوز اضافه نشده‌اند.",
        ),
        (
            "فقط والد",
            "وقتی طول جزو مشخصات محصول است ولی DIN نگفته کدام طول برای کدام قطر مجاز است (مثل DIN 931 و DIN 933)، فرزندی ساخته نمی‌شود و محصول فقط به‌صورت والد می‌ماند. سایزهای موجودش در ستون «سایزهای بدون طول» آمده.",
        ),
        (
            "فرمت کد",
            "DIN912-M8X20 = DIN 912، رزوه M8، طول 20 میلی‌متر | DIN125A-M8 = DIN 125 فرم A برای پیچ M8",
        ),
        ("", ""),
        ("تعداد والد", len(standards)),
        ("   ├ والد با فرزند", len(standards) - n_parent_only),
        ("   └ فقط والد", n_parent_only),
        ("تعداد SKU فرزند", len(skus)),
        ("   ├ با قطر × طول", n_len),
        ("   └ فقط سایز (محصول بدون طول)", len(skus) - n_len),
        ("ترکیب‌های مبهم (برگه Review)", len(review)),
        ("", ""),
        (
            "برگه Review",
            "ترکیب‌هایی که جدول DIN درباره‌شان صریح نیست (سلول «full thread» که روی چند قطر ادغام شده). SKU حساب نشده‌اند.",
        ),
        (
            "اطمینان استخراج",
            "بالا: جدول استاندارد و بدون هشدار | متوسط: سایزها عدد خالی بودند و از برچسب تشخیص داده شدند | پایین: هشدار دارد (ستون هشدارها).",
        ),
        (
            "غیرترجیحی",
            "سایزهای داخل پرانتز در استاندارد، مثل (M14) — مجاز ولی کم‌مصرف.",
        ),
        ("بازتولید", "python3 scripts/extract_skus.py  (ورودی: data/raw/*.xlsx)"),
    ]
    for a, b in lines:
        guide.append([a, b])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 38
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet("Parents", std_cols, standards)
    sheet("SKUs", sku_cols, skus)
    sheet("Review", sku_cols, review)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)


if __name__ == "__main__":
    main()
