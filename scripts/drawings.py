"""Match DIN parents to the REYHER catalogue figures and embed them in a sheet.

data/drawings/reyher/images_index.csv lists one figure per catalogue entry
(standard, title, page, REYHER article codes). A standard can have several
figures, one per form; the form is named in the title as "A = without
chamfer", "ISR = hexalobular socket", ... and is matched to the parent's form
(DIN125A -> A, DIN912TX -> ISR, DIN6797J -> I). "(~similar)" in the catalogue
means the figure shows a close variant, not the exact standard. A figure of a
different form of the same standard is only suggested, never embedded.
"""

import csv
import os
import re

from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR = os.path.join(ROOT, "data", "drawings", "reyher")
INDEX = os.path.join(DIR, "images_index.csv")

MAX_W, MAX_H = 150, 80  # px inside the sheet cell
PAD = 6  # px around the image
CODE = re.compile(r"(?:^|[\s,~])([A-Z][A-Za-z]{0,2}) =")
# parent form suffix -> catalogue form codes
SUFFIX = {"tx": "ISR", "j": "I"}


def load():
    out = {}
    with open(INDEX, encoding="utf-8-sig") as f:
        for i, r in enumerate(csv.DictReader(f)):
            m = re.match(r"DIN (\d+)", r["standard"])
            if not m:
                continue
            r["order"] = i
            r["similar"] = "similar" in r["standard"]
            r["codes"] = set(CODE.findall(r["title"]))
            out.setdefault(m.group(1), []).append(r)
    return out


def wanted(form):
    if not form:
        return set()
    if form in SUFFIX:
        return {SUFFIX[form]}
    if form.endswith("tx"):
        return {form[:-2].upper(), "ISR"}
    if form.endswith("sq"):
        return {form[:-2].upper()}
    return {form.upper()}


def match(standard, catalogue):
    """'DIN 125 a' -> (catalogue row, match label) or (None, '')."""
    m = re.match(r"DIN (\d+)\s*(\S*)", standard)
    cands = catalogue.get(m.group(1)) if m else None
    if not cands:
        return None, ""
    want = wanted(m.group(2).lower())

    def score(r):
        hit = len(want & r["codes"])
        return (-hit, r["similar"], "*" in r["title"], len(r["codes"]), r["order"])

    best = min(cands, key=score)
    hit = want & best["codes"]
    if not want:
        label = "استاندارد"
    elif hit == want:
        label = "فرم دقیق"
    elif hit:
        label = "فرم نزدیک"
    else:
        label = "فرم دیگر"
    if best["similar"]:
        label += " · مشابه"
    return best, label


def annotate(standards):
    """Add figure fields to each parent dict; return number with a figure."""
    catalogue = load()
    n = 0
    for s in standards:
        r, label = match(s["standard"], catalogue)
        other = label.startswith("فرم دیگر")
        use = r if r and not other else None
        s["drawing"] = ""
        s["drawing_file"] = use["image_file"] if use else ""
        s["drawing_title"] = use["title"] if use else ""
        s["reyher_codes"] = use["reyher_codes"] if use else ""
        s["drawing_match"] = (label + " (خالی)") if other else (label or "بدون تصویر")
        s["drawing_suggest"] = f'{r["image_file"]} — {r["title"]}' if other else ""
        n += bool(use)
    return n


def embed(ws, col, file_col):
    """Place the figure named in `file_col` into column `col` of each data row."""
    ws.column_dimensions[get_column_letter(col)].width = (MAX_W + 2 * PAD) / 7
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=file_col).value
        if not name:
            continue
        img = XLImage(os.path.join(DIR, "images", name))
        k = min(MAX_W / img.width, MAX_H / img.height, 1.0)
        w, h = round(img.width * k), round(img.height * k)
        img.width, img.height = w, h
        marker = AnchorMarker(
            col=col - 1,
            colOff=pixels_to_EMU(PAD + (MAX_W - w) // 2),
            row=row - 1,
            rowOff=pixels_to_EMU(PAD + (MAX_H - h) // 2),
        )
        size = XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h))
        img.anchor = OneCellAnchor(_from=marker, ext=size)
        ws.add_image(img)
        ws.row_dimensions[row].height = (MAX_H + 2 * PAD) * 0.75
        for cell in ws[row]:
            cell.alignment = Alignment(vertical="center", wrap_text=cell.alignment.wrap_text)


INDEX_COLS = [  # key, header, width — laid out like a catalogue figure index
    ("standard", "Standard", 14),
    ("drawing", "تصویر", None),
    ("name_en", "Designation (EN)", 38),
    ("drawing_title", "REYHER designation", 46),
    ("parent_code", "Parent code", 13),
    ("drawing_match", "تطبیق تصویر", 20),
    ("drawing_suggest", "تصویر فرم دیگر (پیشنهادی، درج نشده)", 46),
    ("drawing_file", "Figure file", 20),
]


def index_sheet(wb, standards, title="فهرست تصاویر"):
    """Figure index: one row per parent, standard in red beside its figure."""
    from openpyxl.styles import Border, Font, Side

    ws = wb.create_sheet(title)
    line = Side(style="thin", color="9AA3AE")
    box = Border(left=line, right=line, top=line, bottom=line)
    ws.append([c[1] for c in INDEX_COLS])
    for s in standards:
        ws.append([s.get(c[0]) or "" for c in INDEX_COLS])
    for i, c in enumerate(INDEX_COLS, 1):
        if c[2]:
            ws.column_dimensions[get_column_letter(i)].width = c[2]
        h = ws.cell(row=1, column=i)
        h.font = Font(bold=True, color="1F2A44")
        h.border = Border(bottom=Side(style="medium", color="1F2A44"))
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    keys = [c[0] for c in INDEX_COLS]
    embed(ws, keys.index("drawing") + 1, keys.index("drawing_file") + 1)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = box
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].font = Font(bold=True, color="C0392B", size=12)
        row[0].alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws
