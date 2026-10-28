from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.management import call_command
from django.db import DatabaseError, IntegrityError, transaction
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.test import APIClient

from accounting.models import Account, FiscalYear, Period, Voucher
from accounting.services import (
    post_invoice_event,
    post_voucher,
    reverse_voucher,
    trial_balance,
)
from catalog.models import Standard
from catalog.services import create_product, update_product
from core.models import (
    BusinessUnit,
    Capability,
    ModuleFlag,
    OutboxEvent,
    PermissionOverride,
    Role,
    UserScope,
)
from core.services import permissions_for
from crm.models import Party
from sales.models import Document, DocumentLine
from sales.services import issue_document, void_document
from search.parser import normalize, parse
from treasury.services import create_payment


@override_settings(
    CACHES={
        "default": {
            "BACKEND": "django.core.cache.backends.locmem.LocMemCache",
            "LOCATION": "karma-unit-tests",
        }
    }
)
class PlatformTests(TestCase):
    def setUp(self):
        import io

        call_command("seed_platform", verbosity=0, stdout=io.StringIO())
        call_command("seed_accounts", verbosity=0, stdout=io.StringIO())
        self.user = User.objects.create_superuser(
            "owner", "owner@example.test", "a-long-test-only-password"
        )
        self.other = User.objects.create_user("staff", password="a-long-test-only-password")
        self.user.is_verified = lambda: True
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.unit = BusinessUnit.objects.first()
        self.unit2 = BusinessUnit.objects.exclude(pk=self.unit.pk).first()
        self.day = timezone.localdate()
        self.year = FiscalYear.objects.create(
            business_unit=self.unit,
            name="Test",
            start=self.day - timedelta(days=90),
            end=self.day + timedelta(days=90),
        )
        self.period = Period.objects.create(
            year=self.year, start=self.year.start, end=self.year.end
        )
        self.party = Party.objects.create(
            name="مشتری آزمایشی",
            business_unit=self.unit,
            created_by=self.user,
            mobile="09120000000",
        )
        self.standard = Standard.objects.create(code="DIN933")
        self.product = create_product(
            standard=self.standard,
            name="پیچ شش گوش",
            size="M10",
            length="30",
            grade="8.8",
            finish="ZN",
            cost=100,
        )

    def document(self, kind="invoice", **overrides):
        attrs = dict(
            kind=kind,
            business_unit=self.unit,
            customer=self.party,
            created_by=self.user,
            doc_date=self.day,
            tax_rate=10,
            discount=100,
            shipping=50,
        )
        attrs.update(overrides)
        doc = Document.objects.create(**attrs)
        DocumentLine.objects.create(document=doc, product=self.product, quantity=2, unit_price=1000)
        return doc

    def lines(self):
        a = Account.objects.get(business_unit=self.unit, code="1201")
        b = Account.objects.get(business_unit=self.unit, code="6101")
        return [
            {"account": a.pk, "debit": 100, "party": self.party.pk},
            {"account": b.pk, "credit": 100},
        ]

    def post(self, lines=None, key="test"):
        return post_voucher(
            unit_id=self.unit.pk,
            date=self.day,
            actor=self.user,
            source_type="manual",
            source_key=key,
            lines=lines or self.lines(),
        )

    def grant(self, *codes):
        role = Role.objects.create(name="Custom role")
        role.users.add(self.other)
        role.permissions.add(*Capability.objects.filter(code__in=codes))
        UserScope.objects.create(user=self.other, business_unit=self.unit)
        self.client.force_authenticate(self.other)
        return role

    def test_health(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get("/api/v1/health/").status_code, 200)

    def test_dynamic_roles_without_relogin(self):
        r = self.grant("crm.party.view")
        self.assertEqual(self.client.get("/api/v1/parties/").status_code, 200)
        r.permissions.clear()
        self.assertEqual(self.client.get("/api/v1/parties/").status_code, 403)

    def test_deny_override_wins(self):
        self.grant("crm.party.view")
        PermissionOverride.objects.create(
            user=self.other,
            permission=Capability.objects.get(code="crm.party.view"),
            allow=False,
        )
        self.assertNotIn("crm.party.view", permissions_for(self.other))

    def test_scope_direct_object_403(self):
        p = Party.objects.create(name="Other unit", created_by=self.user, business_unit=self.unit2)
        self.grant("crm.party.view")
        self.assertEqual(self.client.get(f"/api/v1/parties/{p.pk}/").status_code, 403)
        self.assertEqual(len(self.client.get("/api/v1/parties/").json()), 1)

    def test_own_only(self):
        self.grant("crm.party.view")
        UserScope.objects.filter(user=self.other).update(own_only=True)
        self.assertEqual(self.client.get("/api/v1/parties/").json(), [])

    def test_cost_hidden_in_api(self):
        self.client.force_authenticate(None)
        self.assertNotIn("cost", self.client.get(f"/api/v1/products/{self.product.pk}/").json())

    def test_cost_hidden_in_excel(self):
        import io

        import openpyxl

        self.client.force_authenticate(None)
        response = self.client.get("/api/v1/products/export/")
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        header = next(wb.active.values)
        self.assertNotIn("قیمت تمام‌شده (ریال)", header)
        self.assertEqual(len(header), 8)

    def test_module_disabled(self):
        ModuleFlag.objects.create(name="catalog", enabled=False)
        self.assertEqual(self.client.get("/api/v1/products/").status_code, 404)

    def test_codes_and_alias_preserved(self):
        code = self.product.system_code
        old = self.product.technical_code
        p = update_product(self.product, length="35")
        self.assertEqual(p.system_code, code)
        result = self.client.get("/api/v1/products/resolve/", {"code": old})
        self.assertEqual(result.json()["id"], p.pk)

    def test_attributes_unique(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            create_product(
                standard=self.standard,
                name="Duplicate",
                size="M10",
                length="30",
                grade="8.8",
                finish="ZN",
            )

    def test_alias_cannot_be_reused(self):
        update_product(self.product, length="35")
        with self.assertRaises(ValidationError):
            create_product(
                standard=self.standard,
                name="Reuse",
                size="M10",
                length="30",
                grade="8.8",
                finish="ZN",
            )

    def test_balanced_posting(self):
        v = self.post()
        self.assertEqual(v.status, "permanent")
        self.assertEqual(sum(x["debit"] - x["credit"] for x in trial_balance(self.unit.pk)), 0)

    def test_unbalanced_rejected(self):
        lines = self.lines()
        lines[1]["credit"] = 99
        with self.assertRaises(ValidationError):
            self.post(lines)
        self.assertFalse(Voucher.objects.exists())

    def test_required_party(self):
        lines = self.lines()
        del lines[0]["party"]
        with self.assertRaises(ValidationError):
            self.post(lines)

    def test_negative_amount(self):
        lines = self.lines()
        lines[0]["debit"] = -100
        with self.assertRaises(ValidationError):
            self.post(lines)

    def test_fractional_rial(self):
        lines = self.lines()
        lines[0]["debit"] = "100.2"
        with self.assertRaises(ValidationError):
            self.post(lines)

    def test_cross_unit_account(self):
        lines = self.lines()
        lines[1]["account"] = Account.objects.get(business_unit=self.unit2, code="6101").pk
        with self.assertRaises(ValidationError):
            self.post(lines)

    def test_closed_period(self):
        self.period.closed = True
        self.period.save()
        with self.assertRaises(ValidationError):
            self.post()

    def test_idempotent_post(self):
        self.assertEqual(self.post().pk, self.post().pk)
        self.assertEqual(Voucher.objects.count(), 1)

    def test_number_rollback(self):
        lines = self.lines()
        lines[1]["credit"] = 99
        with self.assertRaises(ValidationError):
            self.post(lines)
        self.assertEqual(self.post().number, 1)

    def test_database_blocks_permanent_edit(self):
        v = self.post()
        with self.assertRaises(DatabaseError), transaction.atomic():
            Voucher.objects.filter(pk=v.pk).update(source_type="hacked")

    def test_database_blocks_permanent_line_edit(self):
        v = self.post()
        with self.assertRaises(DatabaseError), transaction.atomic():
            v.lines.update(debit=999)

    def test_database_blocks_permanent_line_delete(self):
        v = self.post()
        with self.assertRaises(DatabaseError), transaction.atomic():
            v.lines.all().delete()

    def test_database_blocks_unbalanced_post(self):
        v = Voucher.objects.create(
            business_unit=self.unit,
            period=self.period,
            number=1,
            doc_date=self.day,
            created_by=self.user,
            source_type="bad",
            source_key="bad",
        )
        with self.assertRaises(DatabaseError), transaction.atomic():
            Voucher.objects.filter(pk=v.pk).update(status="permanent")

    def test_reverse_does_not_modify_original(self):
        v = self.post()
        reverse_voucher(v, self.user, self.day)
        v.refresh_from_db()
        self.assertEqual(v.status, "permanent")
        self.assertEqual(sum(x["debit"] - x["credit"] for x in trial_balance(self.unit.pk)), 0)
        for row in trial_balance(self.unit.pk):
            self.assertEqual(row["debit"], row["credit"])

    def test_invoice_totals_snapshot_and_outbox(self):
        doc = issue_document(self.document(), self.user)
        self.assertEqual(doc.total, Decimal(2140))
        self.party.name = "Changed"
        self.party.save()
        self.assertEqual(doc.buyer_snapshot["name"], "مشتری آزمایشی")
        self.assertTrue(
            OutboxEvent.objects.filter(topic="InvoiceIssued", payload={"id": doc.pk}).exists()
        )

    def test_invoice_auto_post_without_opening_accounting(self):
        doc = issue_document(self.document(), self.user)
        from core.tasks import dispatch_outbox

        dispatch_outbox()
        self.assertTrue(Voucher.objects.filter(source_key=f"document:{doc.pk}").exists())
        self.assertEqual(Voucher.objects.count(), 1)
        dispatch_outbox()
        self.assertEqual(Voucher.objects.count(), 1)

    def test_invoice_no_double_issue(self):
        doc = issue_document(self.document(), self.user)
        with self.assertRaises(ValidationError):
            issue_document(doc, self.user)

    def test_invoice_cannot_patch_after_issue(self):
        doc = issue_document(self.document(), self.user)
        self.assertEqual(
            self.client.patch(
                f"/api/v1/documents/{doc.pk}/", {"notes": "hacked"}, format="json"
            ).status_code,
            400,
        )

    def test_backdate_permission(self):
        self.grant("sales.document.issue")
        doc = self.document(doc_date=self.day - timedelta(days=10), created_by=self.other)
        with self.assertRaises(PermissionDenied):
            issue_document(doc, self.other)

    def test_jalali_input(self):
        import jdatetime

        jalali = jdatetime.date.fromgregorian(date=self.day).strftime("%Y-%m-%d")
        response = self.client.post(
            "/api/v1/documents/",
            {
                "business_unit": self.unit.pk,
                "customer": self.party.pk,
                "kind": "invoice",
                "doc_date": jalali,
                "lines": [{"product": self.product.pk, "quantity": "1", "unit_price": "10"}],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Document.objects.get(pk=response.data["id"]).doc_date, self.day)

    def test_void_before_worker_does_not_leave_sales_posted(self):
        doc = issue_document(self.document(), self.user)
        void_document(doc, self.user)
        from core.tasks import dispatch_outbox

        dispatch_outbox()
        self.assertEqual(Voucher.objects.count(), 2)
        for row in trial_balance(self.unit.pk):
            self.assertEqual(row["debit"], row["credit"])

    def test_pdf_is_generated(self):
        doc = issue_document(self.document(), self.user)
        response = self.client.get(f"/api/v1/documents/{doc.pk}/pdf/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF-"))

    def test_payment_and_allocation(self):
        doc = issue_document(self.document(), self.user)
        post_invoice_event(doc.pk)
        bank = Account.objects.get(business_unit=self.unit, code="1102")
        payment = create_payment(
            self.user,
            [{"document": doc.pk, "amount": 1000}],
            business_unit=self.unit,
            party=self.party,
            amount=1000,
            direction="receive",
            account=bank,
            detail={"bank_account": "Test"},
            doc_date=self.day,
        )
        self.assertIsNotNone(payment.voucher)
        self.assertEqual(self.client.get(f"/api/v1/documents/{doc.pk}/").json()["paid"], "1000")

    def test_overallocation_rolls_back(self):
        doc = issue_document(self.document(), self.user)
        bank = Account.objects.get(business_unit=self.unit, code="1102")
        with self.assertRaises(ValidationError):
            create_payment(
                self.user,
                [{"document": doc.pk, "amount": 3000}],
                business_unit=self.unit,
                party=self.party,
                amount=3000,
                direction="receive",
                account=bank,
                detail={"bank_account": "Test"},
                doc_date=self.day,
            )
        from treasury.models import Payment

        self.assertEqual(Payment.objects.count(), 0)

    def test_purchase_uses_periodic_accounts(self):
        doc = issue_document(self.document(kind="purchase"), self.user)
        v = post_invoice_event(doc.pk)
        codes = set(v.lines.values_list("account__code", flat=True))
        self.assertIn("7102", codes)
        self.assertIn("1402", codes)
        self.assertNotIn("1301", codes)

    def test_server_rejects_total_from_client(self):
        self.assertEqual(
            self.client.patch(
                "/api/v1/documents/" + str(self.document().pk) + "/",
                {"total": "1"},
                format="json",
            ).status_code,
            200,
        )
        doc = issue_document(Document.objects.first(), self.user)
        self.assertEqual(doc.total, 2140)

    def test_database_blocks_issued_total_edit(self):
        doc = issue_document(self.document(), self.user)
        with self.assertRaises(DatabaseError), transaction.atomic():
            Document.objects.filter(pk=doc.pk).update(total=1)

    def test_database_blocks_issued_line_add(self):
        doc = issue_document(self.document(), self.user)
        with self.assertRaises(DatabaseError), transaction.atomic():
            DocumentLine.objects.create(
                document=doc, product=self.product, quantity=1, unit_price=1
            )

    def test_login_requires_csrf_and_mfa(self):
        from django.core.cache import cache

        cache.clear()
        client = APIClient(enforce_csrf_checks=True)
        self.assertEqual(
            client.post(
                "/api/v1/session/",
                {"username": "owner", "password": "a-long-test-only-password"},
                format="json",
            ).status_code,
            403,
        )
        token = client.get("/api/v1/session/").json()["csrfToken"]
        response = client.post(
            "/api/v1/session/",
            {"username": "owner", "password": "a-long-test-only-password"},
            format="json",
            HTTP_X_CSRFTOKEN=token,
        )
        self.assertEqual(response.status_code, 403)
        self.assertIn("دومرحله", response.json()["detail"])

    def test_sensitive_session_requires_verified_mfa(self):
        self.user.is_verified = lambda: False
        response = self.client.get("/api/v1/parties/")
        self.assertEqual(response.status_code, 403)

    def test_public_inquiry_daily_limit(self):
        self.client.force_authenticate(None)
        payload = {
            "name": "خریدار",
            "mobile": "۰۹۱۲۳۴۵۶۷۸۹",
            "business_unit": self.unit.pk,
            "items": [{"product": self.product.pk, "quantity": "2"}],
        }
        for _ in range(10):
            self.assertEqual(
                self.client.post("/api/v1/inquiries/", payload, format="json").status_code,
                201,
            )
        self.assertEqual(
            self.client.post("/api/v1/inquiries/", payload, format="json").status_code,
            429,
        )

    def test_fiscal_close_carries_balance_and_details(self):
        from accounting.closing import close_year

        inventory = Account.objects.get(business_unit=self.unit, code="1301")
        retained = Account.objects.get(business_unit=self.unit, code="5301")
        self.post(
            [
                {"account": inventory.pk, "debit": 10000},
                {"account": retained.pk, "credit": 10000},
            ],
            "opening-fixture",
        )
        doc = issue_document(self.document(kind="purchase", discount=0, shipping=0), self.user)
        post_invoice_event(doc.pk)
        next_year = FiscalYear.objects.create(
            business_unit=self.unit,
            name="Next",
            start=self.year.end + timedelta(days=1),
            end=self.year.end + timedelta(days=366),
        )
        next_period = Period.objects.create(
            year=next_year, start=next_year.start, end=next_year.end
        )
        result = close_year(self.year, self.user, ending_inventory=9000, next_period=next_period)
        self.assertEqual(result["cogs"], "3000")
        self.year.refresh_from_db()
        self.assertTrue(self.year.closed)
        close = Voucher.objects.get(pk=result["closing_voucher"])
        opened = Voucher.objects.get(pk=result["opening_voucher"])
        self.assertEqual(
            sum(close.lines.values_list("debit", flat=True)),
            sum(opened.lines.values_list("credit", flat=True)),
        )
        self.assertEqual(opened.lines.get(account__code="1301").debit, 9000)

    def test_payroll_progressive_rates_and_balanced_ledger(self):
        from payroll.models import Contract, PayrollPolicy, Payslip
        from payroll.services import issue_payslip

        contract = Contract.objects.create(
            business_unit=self.unit,
            employee=self.party,
            monthly_salary=100000,
            hourly_overtime=1000,
            created_by=self.user,
        )
        policy = PayrollPolicy.objects.create(
            business_unit=self.unit,
            year=self.year,
            name="Test only; not legal rates",
            rates={
                "employee_insurance_rate": "0.07",
                "employer_insurance_rate": "0.23",
                "insurance_cap": "1000000",
                "tax_brackets": [
                    {"up_to": "50000", "rate": "0"},
                    {"up_to": None, "rate": "0.10"},
                ],
            },
        )
        slip = Payslip.objects.create(
            business_unit=self.unit,
            contract=contract,
            policy=policy,
            doc_date=self.day,
            paid_days=30,
            overtime_hours=5,
            created_by=self.user,
        )
        slip = issue_payslip(slip, self.user)
        self.assertEqual(slip.gross, 105000)
        self.assertEqual(slip.employee_insurance, 7350)
        self.assertEqual(slip.tax, 5500)
        self.assertEqual(slip.net, 92150)
        self.assertEqual(
            sum(slip.voucher.lines.values_list("debit", flat=True)),
            sum(slip.voucher.lines.values_list("credit", flat=True)),
        )

    def test_own_only_ledger_does_not_include_other_users(self):
        self.post()
        self.grant("accounting.voucher.view")
        self.other.is_verified = lambda: True
        UserScope.objects.filter(user=self.other).update(own_only=True)
        self.assertEqual(self.client.get(f"/api/v1/ledger/?unit={self.unit.pk}").json(), [])

    def test_incoming_cheque_all_paths(self):
        from treasury.models import Cheque
        from treasury.services import transition_cheque

        scenarios = [
            ["held", "collecting", "cleared"],
            ["held", "cleared"],
            ["held", "bounced", "held", "collecting", "bounced", "collecting", "cleared"],
            ["held", "collecting", "bounced", "settled"],
            ["held", "endorsed"],
        ]
        for index, states in enumerate(scenarios):
            cheque = Cheque.objects.create(
                business_unit=self.unit,
                party=self.party,
                number=str(index),
                sayad_number=f"{index:016d}",
                bank="test-bank",
                due_date=self.day,
                amount=1000,
                kind="incoming",
                created_by=self.user,
                endorsed_to=self.party,
            )
            for state in states:
                cheque = transition_cheque(cheque, self.user, state, self.day)
                self.assertEqual(cheque.state, state)
                voucher = cheque.transitions.latest("pk").voucher
                self.assertEqual(sum(voucher.lines.values_list("debit", flat=True)), 1000)
                self.assertEqual(sum(voucher.lines.values_list("credit", flat=True)), 1000)

    def test_outgoing_cheque_paths(self):
        from treasury.models import Cheque
        from treasury.services import transition_cheque

        cheque = Cheque.objects.create(
            business_unit=self.unit,
            party=self.party,
            number="out",
            sayad_number="9999999999999999",
            bank="test-bank",
            due_date=self.day,
            amount=1000,
            kind="outgoing",
            created_by=self.user,
        )
        for state in ["issued", "bounced", "issued", "cleared"]:
            cheque = transition_cheque(cheque, self.user, state, self.day)
        with self.assertRaises(ValidationError):
            transition_cheque(cheque, self.user, "issued", self.day)

    def test_ledger_invalid_input_returns_400(self):
        self.assertEqual(self.client.get("/api/v1/ledger/?unit=bad").status_code, 400)
        self.assertEqual(
            self.client.post(
                "/api/v1/ledger/", {"business_unit": self.unit.pk}, format="json"
            ).status_code,
            400,
        )

    def test_pdf_pagination_15_and_40(self):
        import io

        from pypdf import PdfReader

        from documents.services import document_pdf

        for count, expected_pages in [(15, 1), (40, 2)]:
            doc = self.document()
            DocumentLine.objects.bulk_create(
                [
                    DocumentLine(
                        document=doc,
                        product=self.product,
                        quantity=1,
                        unit_price=10,
                        description="ردیف آزمون",
                    )
                    for _ in range(count - 1)
                ]
            )
            doc = issue_document(doc, self.user)
            pdf = PdfReader(io.BytesIO(document_pdf(doc)))
            self.assertEqual(len(pdf.pages), expected_pages)

    def test_reindex_does_not_change_system_codes(self):
        # Import refreshes and search events must not rename existing product identifiers.
        code = self.product.system_code
        update_product(self.product, name="نام اصلاح‌شده")
        self.product.refresh_from_db()
        self.assertEqual(self.product.system_code, code)

    def test_posting_snapshot_survives_rule_change(self):
        from accounting.models import PostingRule

        doc = issue_document(self.document(), self.user)
        PostingRule.objects.filter(business_unit=self.unit, event="invoice").update(lines=[])
        voucher = post_invoice_event(doc.pk)
        self.assertEqual(voucher.lines.count(), 5)

    def test_outbox_failure_does_not_starve_next_invoice(self):
        from core.tasks import dispatch_outbox

        bad = issue_document(self.document(), self.user)
        # Simulate a corrupted/legacy event, rather than disabling a financial assertion.
        OutboxEvent.objects.filter(topic="InvoiceIssued", payload={"id": bad.pk}).update(
            payload={"id": 99999999}
        )
        good = issue_document(self.document(), self.user)
        dispatch_outbox()
        self.assertTrue(Voucher.objects.filter(source_key=f"document:{good.pk}").exists())
        failed = OutboxEvent.objects.get(topic="InvoiceIssued", payload={"id": 99999999})
        self.assertEqual(failed.attempts, 1)
        self.assertIsNotNone(failed.retry_at)

    def test_quote_supply_invoice_packing_chain(self):
        from sales.services import change_quote_state, convert_document, fulfill_supply

        quote = issue_document(self.document(kind="quote"), self.user)
        self.assertEqual(quote.status, "sent")
        quote = change_quote_state(quote, self.user, "approved")
        supply = convert_document(quote, self.user, "supply")
        progress = [
            {
                "line": line.pk,
                "quantity": line.quantity,
                "source": "تأمین‌کننده آزمایشی",
                "notes": "یادداشت چندخطی\nبرای تأمین",
            }
            for line in supply.lines.all()
        ]
        supply = fulfill_supply(supply, self.user, progress, complete=True)
        self.assertEqual(supply.status, "completed")
        invoice = issue_document(convert_document(supply, self.user, "invoice"), self.user)
        packing = issue_document(convert_document(invoice, self.user, "packing"), self.user)
        self.assertEqual(
            len({quote.deal_number, supply.deal_number, invoice.deal_number, packing.deal_number}),
            1,
        )
        self.assertEqual(OutboxEvent.objects.filter(topic="InvoiceIssued").count(), 1)
        quote.refresh_from_db()
        self.assertEqual(quote.status, "invoiced")

    def test_quote_only_permission_cannot_view_invoice(self):
        quote = self.document(kind="quote")
        invoice = self.document()
        self.grant("sales.quote.view")
        self.assertEqual(self.client.get(f"/api/v1/documents/{quote.pk}/").status_code, 200)
        self.assertEqual(self.client.get(f"/api/v1/documents/{invoice.pk}/").status_code, 403)
        rows = self.client.get("/api/v1/documents/").json()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["kind"], "quote")

    def test_returns_cannot_exceed_original(self):
        from sales.services import convert_document

        invoice = issue_document(self.document(), self.user)
        ret = convert_document(invoice, self.user, "return")
        ret.lines.update(quantity=3)
        with self.assertRaises(ValidationError):
            issue_document(ret, self.user)

    def test_multiple_returns_have_unique_numbers(self):
        from sales.services import convert_document

        invoice = issue_document(self.document(discount=0, shipping=0), self.user)
        numbers = []
        for _ in range(2):
            ret = convert_document(invoice, self.user, "return")
            ret.lines.update(quantity=1)
            numbers.append(issue_document(ret, self.user).number)
        self.assertEqual(len(set(numbers)), 2)

    def test_audit_is_append_only_in_database(self):
        from core.models import AuditEntry

        item = AuditEntry.objects.create(actor=self.user, action="audit.test")
        with self.assertRaises(DatabaseError), transaction.atomic():
            AuditEntry.objects.filter(pk=item.pk).update(action="altered")
        with self.assertRaises(DatabaseError), transaction.atomic():
            AuditEntry.objects.filter(pk=item.pk).delete()

    def test_malformed_login_rejected(self):
        client = APIClient()
        self.assertEqual(client.post("/api/v1/session/", [], format="json").status_code, 400)
        self.assertEqual(
            client.post(
                "/api/v1/session/", {"username": [], "password": 1}, format="json"
            ).status_code,
            400,
        )

    def test_sensitive_cost_hidden_until_mfa_verified(self):
        self.user.is_verified = lambda: False
        response = self.client.get(f"/api/v1/products/{self.product.pk}/")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("cost", response.json())

    def test_ledger_api_preserves_large_rial_precision(self):
        value = 9007199254740993
        lines = self.lines()
        lines[0]["debit"] = value
        lines[1]["credit"] = value
        from django.db import connection

        if connection.vendor == "sqlite":
            with self.assertRaises(ValidationError):
                self.post(lines)
            self.assertFalse(Voucher.objects.exists())
            return
        self.post(lines)
        response = self.client.get(f"/api/v1/ledger/?unit={self.unit.pk}").json()
        self.assertIn(str(value), [r["debit"] for r in response])

    def test_json_decimal_literal_never_rounds_via_float(self):
        response = self.client.patch(
            f"/api/v1/products/{self.product.pk}/",
            data='{"cost":9007199254740993e0}',
            content_type="application/json",
        )
        from django.db import connection

        if connection.vendor == "sqlite":
            self.assertEqual(response.status_code, 400, response.data)
            self.product.refresh_from_db()
            self.assertEqual(self.product.cost, 100)
            return
        self.assertEqual(response.status_code, 200, response.data)
        self.product.refresh_from_db()
        self.assertEqual(self.product.cost, Decimal("9007199254740993"))

    def test_oversized_invoice_rolls_back_as_validation_error(self):
        doc = self.document()
        doc.lines.update(
            quantity=Decimal("100000000000000"), unit_price=Decimal("10000000000000000000")
        )
        with self.assertRaises(ValidationError):
            issue_document(doc, self.user)
        doc.refresh_from_db()
        self.assertEqual(doc.status, "draft")

    def test_own_only_cannot_reference_another_users_customer(self):
        self.grant("sales.document.create")
        UserScope.objects.filter(user=self.other).update(own_only=True)
        response = self.client.post(
            "/api/v1/documents/",
            {
                "business_unit": self.unit.pk,
                "customer": self.party.pk,
                "kind": "invoice",
                "doc_date": self.day.isoformat(),
                "lines": [{"product": self.product.pk, "quantity": "1", "unit_price": "100"}],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 403, response.data)
        self.assertFalse(Document.objects.exists())

    def test_draft_cannot_move_across_units_with_old_parent(self):
        parent = self.document(kind="quote")
        child = self.document(parent=parent)
        customer = Party.objects.create(
            name="Other customer", business_unit=self.unit2, created_by=self.user
        )
        response = self.client.patch(
            f"/api/v1/documents/{child.pk}/",
            {"business_unit": self.unit2.pk, "customer": customer.pk},
            format="json",
        )
        self.assertEqual(response.status_code, 400, response.data)

    def test_invoice_cannot_bypass_quote_approval_via_parent_field(self):
        quote = self.document(kind="quote")
        invoice = self.document(parent=quote)
        with self.assertRaises(ValidationError):
            issue_document(invoice, self.user)
        self.assertFalse(OutboxEvent.objects.filter(topic="InvoiceIssued").exists())

    def test_void_payment_releases_allocation_and_preserves_ledger(self):
        from treasury.services import void_payment

        doc = issue_document(self.document(), self.user)
        post_invoice_event(doc.pk)
        bank = Account.objects.get(business_unit=self.unit, code="1102")
        payment = create_payment(
            self.user,
            [{"document": doc.pk, "amount": 1000}],
            business_unit=self.unit,
            party=self.party,
            amount=1000,
            direction="receive",
            account=bank,
            detail={"bank_account": "Test"},
            doc_date=self.day,
        )
        void_payment(payment, self.user, self.day)
        self.assertEqual(self.client.get(f"/api/v1/documents/{doc.pk}/").json()["paid"], "0")
        void_document(doc, self.user)
        for row in trial_balance(self.unit.pk):
            self.assertEqual(row["debit"], row["credit"])

    def test_system_role_cannot_be_deleted_or_stripped(self):
        role = Role.objects.get(name="مدیر سیستم")
        self.assertEqual(
            self.client.patch(
                f"/api/v1/roles/{role.pk}/", {"permissions": []}, format="json"
            ).status_code,
            400,
        )
        self.assertEqual(self.client.delete(f"/api/v1/roles/{role.pk}/").status_code, 400)

    def test_excel_keeps_user_text_literal(self):
        import io

        import openpyxl

        update_product(self.product, name="=1+1")
        response = self.client.get("/api/v1/products/export/")
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        headers = list(next(workbook.active.values))
        column = headers.index("نام کالا") + 1
        cell = workbook.active.cell(2, column)
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_me_only_lists_authorized_business_units(self):
        self.grant("sales.document.view")
        response = self.client.get("/api/v1/me/permissions/").json()
        self.assertEqual([u["id"] for u in response["units"]], [self.unit.pk])


class QueryTests(TestCase):
    def test_document_examples(self):
        cases = [
            (
                "پیچ آلن ۸ در ۲۰ ۱۲.۹",
                {
                    "standards": ["DIN912", "ISO4762"],
                    "size": "M8",
                    "length": "20",
                    "grade": "12.9",
                },
            ),
            (
                "m10x30 8.8 گالوانیزه",
                {"size": "M10", "length": "30", "grade": "8.8", "finish": "ZN"},
            ),
            (
                "شش گوش ۱۲*۵۰",
                {"standards": ["DIN933", "DIN931"], "size": "M12", "length": "50"},
            ),
            ("مهره ۱۶", {"standards": ["DIN934", "ISO4032"], "size": "M16"}),
            ("واشر فنری ۱۰", {"standards": ["DIN127"], "size": "M10"}),
            ("din 933 m 10", {"standards": ["DIN933"], "size": "M10"}),
            ("A325 3/4", {"standards": ["A325"], "size": "3/4"}),
        ]
        for query, expected in cases:
            with self.subTest(query=query):
                result = parse(query)
                for key, value in expected.items():
                    self.assertEqual(result[key], value)

    def test_arabic_normalization(self):
        self.assertEqual(normalize("پيچ ١٢ × ٣٠"), "پیچ 12 X 30")
